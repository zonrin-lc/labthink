#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
refill_template.py — Refill an existing logo-wall PPT template with a new set of logos.

What it does:
  1. opens the template deck (a slide already styled with title + uniform cells)
  2. deletes every PICTURE shape (including pictures nested inside groups)
  3. finds the uniform placeholder cells: AUTO_SHAPEs whose size matches
     --cellw x --cellh inches (default 1.36 x 0.52, tolerance --tol)
  4. sorts cells row-major (rows top-to-bottom, left-to-right within a row)
  5. inserts the logos listed in order.txt, one per cell, fit-inside with
     --pad inner margin, centered; aspect ratio is always preserved
  6. saves to a NEW file — the template itself is never overwritten

Title, cell shapes and decorations are kept untouched. Logo order STRICTLY
follows order.txt (the user's company list order) — never re-sorted.

Requires:  pip install python-pptx pillow

Usage:
  python refill_template.py template.pptx logos_fit order.txt out.pptx \
      --cellw 1.36 --cellh 0.52 --pad 0.06

order.txt: one logo filename (without extension) per line; files are matched
           by prefix, same as build_logo_wall.py.
"""
import os, sys, glob, argparse
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image


def walk(shapes):
    """Yield every non-group shape, descending into groups."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def delete_pictures(shapes):
    """Remove all PICTURE shapes, including those nested in groups."""
    for sh in list(shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            delete_pictures(sh.shapes)
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            sh._element.getparent().remove(sh._element)


def find_cells(slide, cellw_in, cellh_in, tol_in):
    cw, ch, tol = Inches(cellw_in), Inches(cellh_in), Inches(tol_in)
    cells = [sh for sh in walk(slide.shapes)
             if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
             and abs(sh.width - cw) <= tol and abs(sh.height - ch) <= tol]
    # row-major: cluster into rows by top, then sort each row by left
    cells.sort(key=lambda s: (s.top, s.left))
    rows, cur, row_top = [], [], None
    half = Inches(cellh_in) / 2
    for sh in cells:
        if row_top is None or abs(sh.top - row_top) <= half:
            cur.append(sh)
            if row_top is None:
                row_top = sh.top
        else:
            rows.append(sorted(cur, key=lambda s: s.left))
            cur, row_top = [sh], sh.top
    if cur:
        rows.append(sorted(cur, key=lambda s: s.left))
    return [sh for row in rows for sh in row]


def load_order(logodir, orderfile):
    files = {os.path.splitext(os.path.basename(f))[0]: f
             for f in glob.glob(os.path.join(logodir, "*.png"))}
    ordered = []
    with open(orderfile, encoding="utf-8-sig") as f:
        for line in f:
            key = line.strip()
            if not key:
                continue
            if key in files:
                ordered.append(files[key])
                continue
            hit = [v for k, v in files.items() if k.startswith(key) or key in k]
            if hit:
                ordered.append(hit[0])
            else:
                sys.exit(f"ERROR: no logo file matches '{key}'")
    return ordered


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("logodir")
    ap.add_argument("orderfile")
    ap.add_argument("outpptx")
    ap.add_argument("--cellw", type=float, default=1.36, help="cell width, inches")
    ap.add_argument("--cellh", type=float, default=0.52, help="cell height, inches")
    ap.add_argument("--pad", type=float, default=0.06, help="inner margin, inches")
    ap.add_argument("--tol", type=float, default=0.03, help="cell size tolerance, inches")
    args = ap.parse_args()

    prs = Presentation(args.template)
    slide = prs.slides[0]

    delete_pictures(slide.shapes)
    cells = find_cells(slide, args.cellw, args.cellh, args.tol)
    logos = load_order(args.logodir, args.orderfile)

    if len(cells) != len(logos):
        sys.exit(f"ERROR: found {len(cells)} cells but {len(logos)} logos — "
                 f"check --cellw/--cellh/--tol or order.txt")

    max_w, max_h = args.cellw - 2 * args.pad, args.cellh - 2 * args.pad
    for cell, path in zip(cells, logos):
        with Image.open(path) as im:
            w_px, h_px = im.size
        scale = min(max_w / (w_px / 96), max_h / (h_px / 96))  # 96 dpi source
        w_in, h_in = w_px / 96 * scale, h_px / 96 * scale
        left = cell.left + Inches((args.cellw - w_in) / 2)
        top = cell.top + Inches((args.cellh - h_in) / 2)
        slide.shapes.add_picture(path, left, top, Inches(w_in), Inches(h_in))

    prs.save(args.outpptx)
    print(f"saved {args.outpptx}: {len(logos)} logos refilled into template cells")


if __name__ == "__main__":
    main()
