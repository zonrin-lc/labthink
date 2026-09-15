#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_logo_wall.py — Build a uniform client-logo-wall .pptx from normalized logos.

Rules implemented (see references/design-principles.md):
  * uniform cells, logos keep aspect ratio and are centered inside the cell
  * fit-inside: logo <= max_w x max_h, whichever constraint hits first
  * order: STRICTLY the order given in the order file — never re-sorted

Requires:  pip install python-pptx pillow

Usage:
  python build_logo_wall.py logos_fit order.txt out.pptx \
      --title "医药行业服务客户——1000+制药企业的共同选择" \
      --cols 7 --cellw 1.36 --cellh 0.52 --pad 0.06

order.txt: one logo filename (without extension) per line, or "01_国药集团" style;
           the script matches files by prefix.
"""
import os, sys, glob, argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("logodir")
    ap.add_argument("orderfile")
    ap.add_argument("outpptx")
    ap.add_argument("--title", default="")
    ap.add_argument("--cols", type=int, default=7)
    ap.add_argument("--cellw", type=float, default=1.36, help="cell width, inches")
    ap.add_argument("--cellh", type=float, default=0.52, help="cell height, inches")
    ap.add_argument("--pad", type=float, default=0.06, help="inner margin, inches")
    ap.add_argument("--gapx", type=float, default=0.19)
    ap.add_argument("--gapy", type=float, default=0.11)
    args = ap.parse_args()

    files = {os.path.splitext(os.path.basename(f))[0]: f
             for f in glob.glob(os.path.join(args.logodir, "*.png"))}

    ordered = []
    with open(args.orderfile, encoding="utf-8-sig") as f:
        for line in f:
            key = line.strip()
            if not key: continue
            if key in files:
                ordered.append(files[key]); continue
            hit = [v for k, v in files.items()
                   if k.startswith(key) or key in k]
            if hit:
                ordered.append(hit[0])
            else:
                print(f"WARN no file for: {key}", flush=True)
    n = len(ordered)
    if n == 0: sys.exit("no logos matched")

    prs = Presentation()                      # 16:9: 13.333 x 7.5 in
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if args.title:
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.35),
                                      Inches(11.9), Inches(0.7))
        pgh = tb.text_frame.paragraphs[0]
        pgh.text = args.title
        pgh.font.size, pgh.font.bold = Pt(28), True
        pgh.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    cols = args.cols
    rows = (n + cols - 1) // cols
    grid_w = cols * args.cellw + (cols - 1) * args.gapx
    grid_h = rows * args.cellh + (rows - 1) * args.gapy
    x0 = (13.333 - grid_w) / 2
    y0 = 1.3 + (7.5 - 1.3 - 0.4 - grid_h) / 2

    for i, path in enumerate(ordered):
        r, c = divmod(i, cols)
        cx = x0 + c * (args.cellw + args.gapx)
        cy = y0 + r * (args.cellh + args.gapy)
        with Image.open(path) as im:
            w_px, h_px = im.size
        max_w, max_h = args.cellw - 2 * args.pad, args.cellh - 2 * args.pad
        scale = min(max_w / (w_px / 96), max_h / (h_px / 96))  # assume 96 dpi source
        w_in, h_in = w_px / 96 * scale, h_px / 96 * scale
        slide.shapes.add_picture(path, Inches(cx + (args.cellw - w_in) / 2),
                                 Inches(cy + (args.cellh - h_in) / 2),
                                 Inches(w_in), Inches(h_in))
    prs.save(args.outpptx)
    print(f"saved {args.outpptx}: {n} logos, {cols}x{rows} grid")


if __name__ == "__main__":
    main()
